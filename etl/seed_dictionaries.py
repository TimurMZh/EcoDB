"""Seed sites, labs, parameters, aliases, and PDK from Ecology_DB_general.xlsx."""

from __future__ import annotations

import re
from datetime import date
from pathlib import Path

import sys

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
from db import connect  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
DB_INPUT = ROOT / "DB_input"
GENERAL_XLSX = next(DB_INPUT.glob("Ecology_DB_general*.xlsx"))

SITE = {
    "site_code": "OSK",
    "name_ru": "Усть-Каменогорск (Оскемен)",
    "name_en": "Oskemen",
    "region": "Восточно-Казахстанская область",
    "description": "Производственный объект; официальное наименование уточняется.",
}

LABS = [
    ("PEK", "Производственный экологический контроль", "Industrial environmental control"),
    ("ALAB", "ALAB", "ALAB"),
]

# Canonical registry. Codes follow PEK where possible.
PARAMETERS = [
    # air
    ("Dust", "Пыль неорганическая", "Inorganic dust", "mg/m3", "air", 0.01, "air", 10),
    ("CO", "Углерода оксид", "Carbon monoxide", "mg/m3", "air", 0.1, "air", 20),
    ("NO2", "Азота диоксид", "Nitrogen dioxide", "mg/m3", "air", 0.01, "air", 30),
    ("SO2", "Сера диоксид", "Sulfur dioxide", "mg/m3", "air", 0.01, "air", 40),
    ("HCN", "Цианистый водород", "Hydrogen cyanide", "mg/m3", "air", 0.01, "air", 50),
    # water
    ("PH", "рН", "pH", "pH", "surface_water", None, "water", 100),
    ("CN", "Цианиды", "Cyanides", "mg/dm3", "surface_water", 0.01, "water", 110),
    ("CN_FREE", "Цианиды простые", "Free cyanides", "mg/dm3", "surface_water", 0.01, "water", 111),
    ("CN_TOT", "Цианиды общие", "Total cyanides", "mg/dm3", "surface_water", 0.01, "water", 112),
    ("CN_WAD", "Цианиды WAD", "WAD cyanides", "mg/dm3", "surface_water", 0.01, "water", 113),
    ("CNS", "Роданиды", "Thiocyanates", "mg/dm3", "surface_water", 0.01, "water", 114),
    ("SO4", "Сульфаты", "Sulfates", "mg/dm3", "surface_water", 1, "water", 120),
    ("HCO3", "Гидрокарбонаты", "Bicarbonates", "mg/dm3", "surface_water", 1, "water", 130),
    ("Fe", "Железо общее", "Total iron", "mg/dm3", "surface_water", 0.01, "water", 140),
    ("Ca", "Кальций", "Calcium", "mg/dm3", "surface_water", 0.1, "water", 150),
    ("Mg", "Магний", "Magnesium", "mg/dm3", "surface_water", 0.1, "water", 160),
    ("HARDN", "Жесткость общая", "Total hardness", "mg-eq/dm3", "surface_water", None, "water", 170),
    ("As", "Мышьяк", "Arsenic", "mg/dm3", "surface_water", 0.001, "water", 180),
    ("Cl", "Хлориды", "Chlorides", "mg/dm3", "surface_water", 1, "water", 190),
    ("TDS", "Сухой остаток", "Total dissolved solids", "mg/dm3", "surface_water", 1, "water", 200),
    ("SS", "Взвешенные вещества", "Suspended solids", "mg/dm3", "surface_water", 0.1, "water", 210),
    ("COD", "ХПК", "COD", "mg/dm3", "surface_water", 1, "water", 220),
    ("BOD5", "БПК5", "BOD5", "mg/dm3", "surface_water", 0.1, "water", 230),
    ("NO3", "Азот нитратов", "Nitrate nitrogen", "mg/dm3", "surface_water", 0.1, "water", 240),
    ("NH4", "Азот аммонийный", "Ammonium nitrogen", "mg/dm3", "surface_water", 0.05, "water", 250),
    ("PO4", "Фосфаты", "Phosphates", "mg/dm3", "surface_water", 0.01, "water", 260),
    ("SURF", "ПАВ", "Surfactants", "mg/dm3", "surface_water", 0.01, "water", 270),
    ("OIL", "Нефтепродукты", "Petroleum products", "mg/dm3", "surface_water", 0.01, "water", 280),
    ("Cu", "Медь", "Copper", "mg/dm3", "surface_water", 0.001, "water", 300),
    ("Cd", "Кадмий", "Cadmium", "mg/dm3", "surface_water", 0.0001, "water", 310),
    ("Cr", "Хром", "Chromium", "mg/dm3", "surface_water", 0.001, "water", 320),
    ("Mn", "Марганец", "Manganese", "mg/dm3", "surface_water", 0.001, "water", 330),
    ("Ni", "Никель", "Nickel", "mg/dm3", "surface_water", 0.001, "water", 340),
    ("Pb", "Свинец", "Lead", "mg/dm3", "surface_water", 0.001, "water", 350),
    ("Sb", "Сурьма", "Antimony", "mg/dm3", "surface_water", 0.001, "water", 360),
    ("Se", "Селен", "Selenium", "mg/dm3", "surface_water", 0.001, "water", 370),
    ("Zn", "Цинк", "Zinc", "mg/dm3", "surface_water", 0.001, "water", 380),
    ("S_TOT", "Сера общая", "Total sulfur", "mg/dm3", "surface_water", 0.1, "water", 390),
    ("Hg", "Ртуть", "Mercury", "mg/dm3", "surface_water", 0.0001, "water", 400),
    ("Te", "Теллур", "Tellurium", "mg/dm3", "surface_water", 0.001, "water", 410),
    ("DO", "Растворенный кислород", "Dissolved oxygen", "mg/dm3", "surface_water", 0.1, "water", 420),
]

PEK_CODE_NORM = {"Hardn": "HARDN", "Oil prod": "OIL", "Dust": "Dust"}

ALAB_TO_CODE = {
    "рн": "PH",
    "общ.жестк.∑[ca2+mg2+]": "HARDN",
    "са²⁺": "Ca",
    "mg²⁺": "Mg",
    "общ.щелоч. hco₃⁻": "HCO3",
    "cl⁻": "Cl",
    "взвеш. остаток": "SS",
    "сухой остаток": "TDS",
    "cu": "Cu",
    "as (cary)": "As",
    "as": "As",
    "fe (cary)": "Fe",
    "fe": "Fe",
    "cd": "Cd",
    "cr": "Cr",
    "mn": "Mn",
    "ni": "Ni",
    "pb": "Pb",
    "sb": "Sb",
    "se": "Se",
    "zn": "Zn",
    "s общ.": "S_TOT",
    "s сульфат": "SO4",
    "сn⁻ пр.": "CN_FREE",
    "cn- общий": "CN_TOT",
    "cn- wad.": "CN_WAD",
    "cns⁻": "CNS",
}

PDK_WATER_TO_CODE = {
    "pH": "PH",
    "Цианиды": "CN",
    "Сульфаты": "SO4",
    "Железо общее": "Fe",
    "Жесткость общая": "HARDN",
    "Хлориды": "Cl",
    "Сухой остаток": "TDS",
    "Мышьяк": "As",
}

PDK_AIR_TO_CODE = {
    "Пыль неорганическая": "Dust",
    "Углерода оксид": "CO",
    "Азота диоксид": "NO2",
    "Сера диоксид": "SO2",
}

EXTRA_ALIASES = [
    ("Dust", "пыль неорганическая 70-20%", "PDF"),
    ("Dust", "пыль неорганическая 70-20% ", "PDF"),
    ("HCN", "цианид водорода", "PDF"),
    ("HCN", "цианистый водород", "PEK"),
    ("CO", "углерода оксид", "PEK"),
    ("NO2", "азота диоксид", "PEK"),
    ("SO2", "сера диоксид", "PEK"),
    ("Fe", "железо", None),
    ("PH", "рн (водородный показатель)", None),
    ("Hg", "ртуть", None),
    ("Te", "теллур", None),
    ("DO", "растворенный кислород", None),
]

# Anonymous report numbers plus named wells from the PDF (table 9).
WATER_POINTS = [
    *(
        (
            f"SW-{i}",
            "surface_water",
            "створ",
            f"Поверхностные воды, точка {i}",
            "Номер группы в «Результаты мониторинга поверхностных вод». "
            "Имя створа и координаты в файле не указаны.",
        )
        for i in range(1, 5)
    ),
    *(
        (
            f"GW-{i}",
            "groundwater",
            "точка отчёта",
            f"Подземные воды, точка {i}",
            "Номер группы в «Результаты мониторинга подземных вод». Координаты не указаны."
            + (
                f" Значения 1 кв. 2025 совпадают с SW-{i}; до подтверждения заказчика это отдельная точка."
                if i <= 4
                else ""
            ),
        )
        for i in range(1, 11)
    ),
    (
        "ФС-1",
        "groundwater",
        "скважина",
        "Скважина ФС-1",
        "PDF, таблица 9. Координаты в колонке пустые.",
    ),
    (
        "ФС-4",
        "groundwater",
        "скважина",
        "Скважина ФС-4",
        "PDF, таблица 9. Координаты в колонке пустые.",
    ),
    (
        "ФС-5",
        "groundwater",
        "скважина",
        "Скважина ФС-5",
        "PDF, таблица 9. Координаты в колонке пустые.",
    ),
    (
        "НС-2",
        "groundwater",
        "скважина",
        "Скважина НС-2",
        "PDF, таблица 9 обрывается на кальции. Координаты пустые.",
    ),
]


def norm_label(text: str) -> str:
    text = text.replace("ё", "е").replace("Ё", "е")
    text = re.sub(r"\s+", " ", text.strip().lower())
    return text


def _fetch_id(cur, sql: str, params) -> int:
    cur.execute(sql, params)
    row = cur.fetchone()
    if not row:
        raise RuntimeError(f"expected id from {sql!r}")
    return row[0]


def upsert_alias(cur, table: str, id_col: str, id_value: int, alias: str, source: str | None):
    alias_n = norm_label(alias)
    if not alias_n:
        return
    cur.execute(
        f"""
        INSERT INTO eco.{table} ({id_col}, alias_norm, source_system)
        SELECT %s, %s, %s
        WHERE NOT EXISTS (
            SELECT 1 FROM eco.{table}
            WHERE alias_norm = %s
              AND COALESCE(source_system, '') = COALESCE(%s, '')
        )
        """,
        (id_value, alias_n, source, alias_n, source),
    )


def _no_pdk(val) -> bool:
    if val is None:
        return True
    if isinstance(val, str) and "нет" in val.lower():
        return True
    return False


def collapse_year_limits(years: list[int], values: list) -> list[tuple[int, int, float]]:
    runs: list[tuple[int, int, float]] = []
    start = prev_year = None
    prev_val = None
    for year, val in zip(years, values):
        if _no_pdk(val):
            if start is not None:
                runs.append((start, prev_year, float(prev_val)))
                start = None
            continue
        num = float(val)
        if start is None:
            start = prev_year = year
            prev_val = num
        elif num == prev_val:
            prev_year = year
        else:
            runs.append((start, prev_year, float(prev_val)))
            start = prev_year = year
            prev_val = num
    if start is not None:
        runs.append((start, prev_year, float(prev_val)))
    return runs


def seed() -> None:
    wb = load_workbook(GENERAL_XLSX, data_only=True, read_only=True)

    conn = connect()
    conn.autocommit = False
    try:
        cur = conn.cursor()

        cur.execute(
            """
            INSERT INTO eco.sites (site_code, name_ru, name_en, region, description)
            VALUES (%(site_code)s, %(name_ru)s, %(name_en)s, %(region)s, %(description)s)
            ON CONFLICT (site_code) DO UPDATE
            SET name_ru = EXCLUDED.name_ru,
                name_en = EXCLUDED.name_en,
                region = EXCLUDED.region,
                description = EXCLUDED.description
            RETURNING site_id
            """,
            SITE,
        )
        site_id = cur.fetchone()[0]

        lab_ids = {}
        for code, name_ru, name_en in LABS:
            cur.execute(
                """
                INSERT INTO eco.laboratories (lab_code, name_ru, name_en)
                VALUES (%s, %s, %s)
                ON CONFLICT (lab_code) DO UPDATE
                SET name_ru = EXCLUDED.name_ru, name_en = EXCLUDED.name_en
                RETURNING lab_id
                """,
                (code, name_ru, name_en),
            )
            lab_ids[code] = cur.fetchone()[0]

        param_ids = {}
        for row in PARAMETERS:
            code, name_ru, name_en, unit, media, mdl, group, sort = row
            cur.execute(
                """
                INSERT INTO eco.parameters (
                    param_code, name_ru, name_en, unit, media_type,
                    mdl_default, group_name, sort_order
                )
                VALUES (%s, %s, %s, %s, %s, %s, %s, %s)
                ON CONFLICT (param_code) DO UPDATE
                SET name_ru = EXCLUDED.name_ru,
                    name_en = EXCLUDED.name_en,
                    unit = EXCLUDED.unit,
                    media_type = EXCLUDED.media_type,
                    mdl_default = EXCLUDED.mdl_default,
                    group_name = EXCLUDED.group_name,
                    sort_order = EXCLUDED.sort_order
                RETURNING param_id
                """,
                (code, name_ru, name_en, unit, media, mdl, group, sort),
            )
            param_ids[code] = cur.fetchone()[0]
            upsert_alias(cur, "parameter_aliases", "param_id", param_ids[code], name_ru, None)
            if name_en:
                upsert_alias(cur, "parameter_aliases", "param_id", param_ids[code], name_en, None)
            upsert_alias(cur, "parameter_aliases", "param_id", param_ids[code], code, None)

        for i in range(1, 31):
            code = f"Т-{i}"
            cur.execute(
                """
                INSERT INTO eco.monitoring_points (
                    point_code, site_id, name_ru, name_en, media_type, point_type, comments
                )
                VALUES (%s, %s, %s, %s, 'air', 'СЗЗ', %s)
                ON CONFLICT (point_code) DO UPDATE
                SET site_id = EXCLUDED.site_id,
                    name_ru = EXCLUDED.name_ru,
                    media_type = EXCLUDED.media_type,
                    point_type = EXCLUDED.point_type
                RETURNING point_id
                """,
                (
                    code,
                    site_id,
                    f"Точка {code} (граница СЗЗ)",
                    f"Point {code} (SPZ boundary)",
                    "Координаты не переданы; geom NULL до каталога точек.",
                ),
            )
            point_id = cur.fetchone()[0]
            for alias in (code, f"T-{i}", f"Т.№{i}", f"Т.№ {i}", f"т.№{i}"):
                upsert_alias(cur, "point_aliases", "point_id", point_id, alias, None)
            upsert_alias(cur, "point_aliases", "point_id", point_id, f"Т.№{i}", "PDF")

        for code, media, point_type, name_ru, comments in WATER_POINTS:
            cur.execute(
                """
                INSERT INTO eco.monitoring_points (
                    point_code, site_id, name_ru, media_type, point_type, comments
                )
                VALUES (%s, %s, %s, %s, %s, %s)
                ON CONFLICT (point_code) DO UPDATE
                SET site_id = EXCLUDED.site_id,
                    name_ru = EXCLUDED.name_ru,
                    media_type = EXCLUDED.media_type,
                    point_type = EXCLUDED.point_type,
                    comments = EXCLUDED.comments
                RETURNING point_id
                """,
                (code, site_id, name_ru, media, point_type, comments),
            )
            point_id = cur.fetchone()[0]
            upsert_alias(cur, "point_aliases", "point_id", point_id, code, None)
            upsert_alias(cur, "point_aliases", "point_id", point_id, name_ru, "PDF")

        # PEK / ALAB synonym sheets
        ws = wb["Показатели_вода_ПЭК"]
        for row in ws.iter_rows(values_only=True):
            if not row or not row[0] or not row[1]:
                continue
            raw_name, raw_code = str(row[0]), str(row[1]).strip()
            code = PEK_CODE_NORM.get(raw_code, raw_code)
            if code not in param_ids:
                continue
            upsert_alias(cur, "parameter_aliases", "param_id", param_ids[code], raw_name, "PEK")
            upsert_alias(cur, "parameter_aliases", "param_id", param_ids[code], raw_code, "PEK")

        ws = wb["Показатели_воздух_ПЭК"]
        for row in ws.iter_rows(values_only=True):
            if not row or not row[0] or not row[1]:
                continue
            raw_name, raw_code = str(row[0]), str(row[1]).strip()
            code = PEK_CODE_NORM.get(raw_code, raw_code)
            if code not in param_ids:
                continue
            upsert_alias(cur, "parameter_aliases", "param_id", param_ids[code], raw_name, "PEK")
            upsert_alias(cur, "parameter_aliases", "param_id", param_ids[code], raw_code, "PEK")
            # first line of multiline cell = name used in air Excel
            upsert_alias(
                cur, "parameter_aliases", "param_id", param_ids[code], raw_name.split("\n")[0], "PEK"
            )

        ws = wb["Показатели_вода_ALAB"]
        unmatched_alab = []
        for row in ws.iter_rows(values_only=True):
            if not row or not row[0]:
                continue
            raw_name = str(row[0])
            key = norm_label(raw_name)
            code = ALAB_TO_CODE.get(key)
            if not code or code not in param_ids:
                unmatched_alab.append(raw_name.strip())
                continue
            upsert_alias(cur, "parameter_aliases", "param_id", param_ids[code], raw_name, "ALAB")
        if unmatched_alab:
            print("Unmapped ALAB names:", unmatched_alab)

        for code, alias, src in EXTRA_ALIASES:
            upsert_alias(cur, "parameter_aliases", "param_id", param_ids[code], alias, src)

        for std_code, name_ru, name_en, media in (
            ("PDK_MR_AIR", "ПДК максимально разовая, воздух", "Ambient air MPC (max-single)", "air"),
            ("PDK_WATER", "ПДК вода", "Water MPC", "surface_water"),
        ):
            cur.execute(
                """
                INSERT INTO eco.guideline_standards (standard_code, name_ru, name_en, media_type, authority)
                VALUES (%s, %s, %s, %s, 'Ecology_DB_general')
                ON CONFLICT (standard_code) DO UPDATE
                SET name_ru = EXCLUDED.name_ru, name_en = EXCLUDED.name_en
                RETURNING standard_id
                """,
                (std_code, name_ru, name_en, media),
            )
        air_std = _fetch_id(
            cur, "SELECT standard_id FROM eco.guideline_standards WHERE standard_code = %s", ("PDK_MR_AIR",)
        )
        water_std = _fetch_id(
            cur, "SELECT standard_id FROM eco.guideline_standards WHERE standard_code = %s", ("PDK_WATER",)
        )

        def seed_pdk(sheet_name: str, name_map: dict[str, str], standard_id: int, media: str, unit: str):
            ws = wb[sheet_name]
            rows = list(ws.iter_rows(values_only=True))
            header = rows[0]
            years = [int(y) for y in header[1:] if y is not None]
            for row in rows[1:]:
                name = row[0]
                if not name:
                    continue
                code = name_map.get(str(name).strip())
                if not code or code not in param_ids:
                    raise RuntimeError(f"no param mapping for PDK '{name}'")
                for y0, y1, limit in collapse_year_limits(years, list(row[1 : 1 + len(years)])):
                    valid_from = date(y0, 1, 1)
                    valid_to = None if y1 == years[-1] else date(y1, 12, 31)
                    cur.execute(
                        """
                        INSERT INTO eco.guideline_values (
                            standard_id, param_id, media_type, comparison_op,
                            limit_value, unit, valid_from, valid_to, comments
                        )
                        VALUES (%s, %s, %s, 'GT', %s, %s, %s, %s, %s)
                        ON CONFLICT (standard_id, param_id, media_type, valid_from) DO UPDATE
                        SET limit_value = EXCLUDED.limit_value,
                            valid_to = EXCLUDED.valid_to,
                            unit = EXCLUDED.unit,
                            comments = EXCLUDED.comments
                        """,
                        (
                            standard_id,
                            param_ids[code],
                            media,
                            limit,
                            unit,
                            valid_from,
                            valid_to,
                            f"{name}: {y0}–{y1} из {GENERAL_XLSX.name}",
                        ),
                    )

        seed_pdk("ПДК воздух", PDK_AIR_TO_CODE, air_std, "air", "mg/m3")
        seed_pdk("ПДК вода", PDK_WATER_TO_CODE, water_std, "surface_water", "mg/dm3")

        # HCN is in air results/PDF but missing from ПДК воздух sheet.
        cur.execute(
            """
            INSERT INTO eco.guideline_values (
                standard_id, param_id, media_type, comparison_op,
                limit_value, unit, valid_from, valid_to, comments
            )
            VALUES (%s, %s, 'air', 'GT', 0.01, 'mg/m3', DATE '2020-01-01', NULL,
                    'ПДК 0.01 мг/м³ по протоколам и колонке PDK файла воздуха (не из листа «ПДК воздух»)')
            ON CONFLICT (standard_id, param_id, media_type, valid_from) DO UPDATE
            SET limit_value = EXCLUDED.limit_value, comments = EXCLUDED.comments
            """,
            (air_std, param_ids["HCN"]),
        )

        conn.commit()
        print(
            f"Seeded site OSK, {len(lab_ids)} labs, {len(param_ids)} parameters, "
            f"30 air points, {len(WATER_POINTS)} water points, PDK from {GENERAL_XLSX.name}"
        )
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()
        wb.close()


if __name__ == "__main__":
    seed()
