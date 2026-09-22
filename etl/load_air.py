"""Load air Excel into staging_lab_import and process batch AIR_PILOT."""

from __future__ import annotations

import json
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from psycopg2.extras import Json, execute_values

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(Path(__file__).resolve().parent))

from db import connect  # noqa: E402

AIR_XLSX = ROOT / "DB_input" / "Результаты мониторинга воздуха.xlsx"
BATCH = "AIR_PILOT"
LAB_CODE = "PEK"
HCN_LABELS = {"цианистый водород", "hcn"}


def _cell_text(val) -> str | None:
    if val is None:
        return None
    if isinstance(val, datetime):
        return val.date().isoformat()
    if isinstance(val, date):
        return val.isoformat()
    if isinstance(val, float):
        if val == int(val) and abs(val) < 1e12:
            return str(int(val))
        return format(val, "g")
    text = str(val).strip()
    return text or None


def _is_hcn(name: str) -> bool:
    n = " ".join(name.lower().replace("ё", "е").split())
    return n in HCN_LABELS or n.startswith("цианист")


def clear_air_pilot(cur) -> None:
    cur.execute(
        """
        DELETE FROM eco.results r
        USING eco.samples sm
            JOIN eco.sampling_events e ON e.event_id = sm.event_id
            JOIN eco.monitoring_points p ON p.point_id = e.point_id
            JOIN eco.sites s ON s.site_id = p.site_id
        WHERE r.sample_id = sm.sample_id
          AND s.site_code = 'OSK'
          AND e.media_type = 'air'
        """
    )
    cur.execute(
        """
        DELETE FROM eco.samples sm
        USING eco.sampling_events e
            JOIN eco.monitoring_points p ON p.point_id = e.point_id
            JOIN eco.sites s ON s.site_id = p.site_id
        WHERE sm.event_id = e.event_id
          AND s.site_code = 'OSK'
          AND e.media_type = 'air'
        """
    )
    cur.execute(
        """
        DELETE FROM eco.sampling_events e
        USING eco.monitoring_points p
            JOIN eco.sites s ON s.site_id = p.site_id
        WHERE e.point_id = p.point_id
          AND s.site_code = 'OSK'
          AND e.media_type = 'air'
        """
    )
    cur.execute(
        "DELETE FROM eco.staging_lab_import WHERE batch_code = %s",
        (BATCH,),
    )


def load_air(*, reload: bool = True, impute_hcn: bool = True) -> None:
    if not AIR_XLSX.exists():
        raise FileNotFoundError(AIR_XLSX)

    wb = load_workbook(AIR_XLSX, data_only=True, read_only=True)
    ws = wb.active
    rows = []
    empty_kept = 0
    hcn_imputed = 0

    for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if i == 1:
            continue
        ind, dt, pt, val, pdk, year, quarter = (list(row) + [None] * 7)[:7]
        if ind is None and pt is None:
            continue

        extra = {"source_pdk": _cell_text(pdk)}
        result_text = _cell_text(val)
        if result_text is None and impute_hcn and ind and _is_hcn(str(ind)):
            result_text = "не обн."
            extra["imputed_nondetect"] = True
            extra["source_value"] = None
            hcn_imputed += 1
        elif result_text is None:
            empty_kept += 1

        coll_date = None
        if isinstance(dt, datetime):
            coll_date = dt.date()
        elif isinstance(dt, date):
            coll_date = dt

        rows.append(
            (
                BATCH,
                AIR_XLSX.name,
                ws.title,
                i,
                LAB_CODE,
                None if pt is None else str(pt).strip(),
                "air",
                None if ind is None else str(ind).strip(),
                coll_date,
                result_text,
                _cell_text(pdk),
                int(year) if year is not None else None,
                int(quarter) if quarter is not None else None,
                Json(extra),
            )
        )
    wb.close()

    conn = connect()
    try:
        cur = conn.cursor()
        if reload:
            clear_air_pilot(cur)

        execute_values(
            cur,
            """
            INSERT INTO eco.staging_lab_import (
                batch_code, source_file, source_sheet, row_no,
                lab_code, point_code, media_type, param_name_raw,
                collection_date, result_text, pdk_raw,
                year_num, quarter_num, extra_json
            ) VALUES %s
            """,
            rows,
            page_size=500,
        )
        inserted = len(rows)
        conn.commit()

        cur.execute("SELECT metric, value FROM eco.fn_process_staging_import(%s)", (BATCH,))
        metrics = dict(cur.fetchall())
        conn.commit()

        cur.execute(
            """
            SELECT COALESCE(process_error, '(ok)'), count(*)
            FROM eco.staging_lab_import
            WHERE batch_code = %s
            GROUP BY 1
            ORDER BY 2 DESC
            """,
            (BATCH,),
        )
        errors = cur.fetchall()

        cur.execute(
            """
            SELECT count(*) FROM eco.results r
            JOIN eco.samples sm ON sm.sample_id = r.sample_id
            JOIN eco.sampling_events e ON e.event_id = sm.event_id
            WHERE e.media_type = 'air'
            """
        )
        n_results = cur.fetchone()[0]
        cur.execute("SELECT count(*) FROM eco.result_exceedances")
        n_exc = cur.fetchone()[0]
        cur.execute(
            """
            SELECT p.point_code, pr.param_code, e.event_date, r.result_value, x.ratio
            FROM eco.result_exceedances x
            JOIN eco.results r ON r.result_id = x.result_id
            JOIN eco.parameters pr ON pr.param_id = r.param_id
            JOIN eco.samples sm ON sm.sample_id = r.sample_id
            JOIN eco.sampling_events e ON e.event_id = sm.event_id
            JOIN eco.monitoring_points p ON p.point_id = e.point_id
            ORDER BY x.ratio DESC
            LIMIT 10
            """
        )
        top = cur.fetchall()

        print(f"AIR_PILOT staging rows: {inserted}")
        print(f"HCN empty→«не обн.»: {hcn_imputed}; other empty kept: {empty_kept}")
        print("fn_process_staging_import:", json.dumps(metrics, ensure_ascii=False))
        print("staging status:")
        for msg, cnt in errors:
            print(f"  {cnt:5d}  {msg}")
        print(f"air results: {n_results}; exceedances: {n_exc}")
        if top:
            print("top exceedances:")
            for rec in top:
                print(f"  {rec}")
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


if __name__ == "__main__":
    reload = "--no-reload" not in sys.argv
    impute = "--no-impute-hcn" not in sys.argv
    load_air(reload=reload, impute_hcn=impute)
